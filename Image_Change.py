from tkinter import Tk, filedialog, Button, Label, Frame, StringVar, OptionMenu
from aip import AipOcr
import os
from docx import Document
import pdfkit
import tkinter.messagebox

# 百度智能云的API配置信息
APP_ID = '117467950'
API_KEY = 'jUAsIg4Lzmtyk5P6R3oBP7vb'
SECRET_KEY = 'XrWPWexiOJ1R4wmsAvo9gYMaMA4WqmlX'


def api_connect(image):
    client = AipOcr(APP_ID, API_KEY, SECRET_KEY)
    options = {"detect_direction": "true"}
    result = client.basicGeneral(image, options)
    text = ""
    words_result = result.get("words_result", [])
    for word in words_result:
        text += word['words'] + '\n'
    return text


def select_image():
    Tk().withdraw()
    image_path = filedialog.askopenfilename(
        title="选择图片文件",
        filetypes=[("Image Files", "*.jpg *.png *.jpeg *.bmp *.tiff")]
    )
    return image_path


def save_document(text, save_format):
    Tk().withdraw()
    if not text:
        tkinter.messagebox.showerror("错误", "没有可保存的内容，请先进行转换。")
        return

    if save_format == 'Word Document':
        save_path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx")],
            title="保存文档"
        )
        if save_path:
            doc = Document()
            for row in text:
                table = doc.add_table(rows=1, cols=len(row))
                table.style = 'Table Grid'
                for i, cell in enumerate(row):
                    table.cell(0, i).text = cell
            doc.save(save_path)
    elif save_format == 'Text Files':
        save_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text Files", "*.txt")],
            title="保存文档"
        )
        if save_path:
            with open(save_path, "w", encoding="utf-8") as fp:
                fp.write('\n'.join(['\t'.join(row) for row in text]))
    elif save_format == 'PDF Document':
        save_path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF Files", "*.pdf")],
            title="保存文档"
        )
        if save_path:
            # 指定wkhtmltopdf的路径
            config = pdfkit.configuration(wkhtmltopdf=r'D:\pdf_tool\wkhtmltopdf\bin\wkhtmltopdf.exe')
            # 使用UTF-8编码
            options = {'encoding': 'UTF-8'}
            pdfkit.from_string('\n'.join(['\t'.join(row) for row in text]), save_path, configuration=config,
                               options=options)


def process_image(image_path):
    with open(image_path, 'rb') as f:
        image = f.read()
    text = api_connect(image)
    # 假设识别结果是一个字符串，需要进一步处理成表格格式
    # 这里需要根据实际情况进行处理，例如按行分割，再按列分割
    lines = text.strip().split('\n')
    table = []
    for line in lines:
        cells = line.split('\t')
        table.append(cells)
    return table


def main():
    root = Tk()
    root.title("图片处理工具")
    root.geometry("400x400")
    root.configure(bg="#f0f0f0")

    # 创建框架
    frame = Frame(root, bg="#f0f0f0")
    frame.pack(expand=True, fill="both", padx=20, pady=20)

    # 添加标签
    label = Label(frame, text="请选择图片文件进行处理", font=("Arial", 12), bg="#f0f0f0")
    label.pack(pady=10)

    selected_image_path = StringVar()
    document_text = None
    save_format = StringVar(value="Word Document")

    # 添加文件格式选择下拉菜单
    format_label = Label(frame, text="选择保存格式：", font=("Arial", 10), bg="#f0f0f0")
    format_label.pack(pady=5)
    format_options = ["Word Document", "Text Files", "PDF Document"]
    format_menu = OptionMenu(frame, save_format, *format_options)
    format_menu.pack(pady=5)

    # 添加按钮
    def start_conversion():
        nonlocal document_text
        image_path = select_image()
        if image_path:
            selected_image_path.set(image_path)
            document_text = process_image(image_path)
            button_save.config(state="normal")

    button_select = Button(frame, text="开始转换", command=start_conversion, width=15, bg="#4CAF50", fg="white",
                           font=("Arial", 10))
    button_select.pack(pady=5)

    button_save = Button(frame, text="保存文档", command=lambda: save_document(document_text, save_format.get()),
                         width=15, bg="#4CAF50", fg="white", font=("Arial", 10))
    button_save.pack(pady=5)
    button_save.config(state="disabled")

    # 添加退出按钮
    button_exit = Button(frame, text="退出", command=root.quit, width=15, bg="#f44336", fg="white", font=("Arial", 10))
    button_exit.pack(pady=5, side="left", anchor="w")

    root.mainloop()


if __name__ == "__main__":
    main()
